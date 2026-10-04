"""Import the recovered export as historical evidence, never as a live crawl."""

from contextlib import closing
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from phone_assistant.catalog import CATALOG_SOURCE
from phone_assistant.knowledge import PROJECT_ROOT
from phone_assistant.storage import Storage


def _cell_text(value) -> str:
    return value.isoformat() if isinstance(value, (datetime, date)) else str(value)


def legacy_rows(root: Path = PROJECT_ROOT):
    path = Path(root) / CATALOG_SOURCE
    if not path.is_file():
        raise FileNotFoundError(f"缺少原始历史手机Excel：{path}")
    with closing(load_workbook(path, read_only=True, data_only=True)) as workbook:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows)
        for row_number, row in enumerate(rows, 2):
            if all(value is None or isinstance(value, str) and not value.strip() for value in row):
                continue
            record = dict(zip(headers, row))
            id = str(record["产品ID"]) if record["产品ID"] is not None else None
            yield {
                "id": id, "name": record["产品型号"], "brand": record["品牌"],
                "price": record["参考价格(人民币)"], "image_url": None,
                "source_url": None, "fetched_at": None,
                "availability": "historical", "origin": "legacy",
                "legacy_file": CATALOG_SOURCE, "legacy_row": row_number,
                "legacy_source": f"{CATALOG_SOURCE}#row={row_number}",
                "specs": {key: _cell_text(value) if value is not None else None for key, value in record.items()},
            }


def import_legacy(storage: Storage, root: Path = PROJECT_ROOT) -> dict:
    return storage.import_many(legacy_rows(root))
