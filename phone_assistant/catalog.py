"""Browse the recovered spreadsheet while preserving its original cell values."""

from contextlib import closing
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

from phone_assistant.knowledge import PROJECT_ROOT


CATALOG_SOURCE = "data/recovered/phone_specs_export_中文表头.xlsx"
_BRAND_ALIASES = {
    "xiaomi": "小米",
    "redmi": "红米",
    "huawei": "华为",
    "honor": "荣耀",
    "oneplus": "一加",
    "apple": "苹果",
    "iphone": "苹果",
    "samsung": "三星",
    "google": "谷歌",
    "sony": "索尼",
    "motorola": "摩托罗拉",
    "meizu": "魅族",
    "nokia": "诺基亚",
    "philips": "飞利浦",
    "lenovo": "联想",
    "zte": "中兴",
    "asus": "华硕",
    "blackberry": "黑莓",
}


def _canonical(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).casefold()
    for english, chinese in _BRAND_ALIASES.items():
        text = re.sub(rf"(?<![a-z]){english}", chinese, text)
    text = text.replace("真我", "realme")
    return re.sub(r"苹果\s*苹果", "苹果", text)


def _compact(text: str) -> str:
    return re.sub(r"[\s_()\-]", "", _canonical(text))


def _blank(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _year(record: dict) -> int | None:
    value = record["发布年份"]
    return None if _blank(value) else int(value)


def _model_aliases(record: dict) -> set[str]:
    if _blank(record["产品型号"]):
        return set()
    full_name = _canonical(str(record["产品型号"]))
    base_name = re.sub(r"\([^()]*(?:gb|tb|mb)[^()]*\)\s*$", "", full_name)
    names = {_compact(full_name), _compact(base_name)}
    if not _blank(record["品牌"]):
        brand = _compact(str(record["品牌"]))
        names.update(brand + name for name in tuple(names) if not name.startswith(brand))
    if not _blank(record["系列"]):
        series = _compact(str(record["系列"]))
        # A distinctive family name such as Find X8 or Pura 70 can omit the brand.
        if len(series) >= 4 and re.match(r"[a-z]", series) and re.search(r"\d", series):
            names.add(series)
    return names


class PhoneCatalog:
    def __init__(self, root: Path = PROJECT_ROOT):
        self.source = CATALOG_SOURCE
        self.path = Path(root) / CATALOG_SOURCE
        if not self.path.is_file():
            raise FileNotFoundError(f"缺少手机数据表：{self.path}")
        with closing(load_workbook(self.path, read_only=True, data_only=True)) as workbook:
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows)
            self.records = [
                dict(zip(headers, row)) for row in rows if any(not _blank(value) for value in row)
            ]
        self.brands = sorted({record["品牌"] for record in self.records if not _blank(record["品牌"])})
        self._years = [_year(record) for record in self.records]
        self._brands = [
            _compact(str(record["品牌"])) if not _blank(record["品牌"]) else ""
            for record in self.records
        ]
        self._search_text = [
            _compact(" ".join(str(record[field]) for field in ("产品型号", "品牌", "系列") if not _blank(record[field])))
            for record in self.records
        ]
        aliases: dict[str, list[int]] = {}
        for index, record in enumerate(self.records):
            for name in _model_aliases(record):
                aliases.setdefault(name, []).append(index)
        self._model_patterns = [
            (
                re.compile(
                    r"(?<![a-z0-9])"
                    + r"[\s_-]*".join(re.escape(character) for character in name)
                    + r"(?=$|[^a-z0-9+])"
                ),
                indexes,
            )
            for name, indexes in aliases.items()
        ]

    @property
    def stats(self) -> dict[str, int | None]:
        years = [year for year in self._years if year is not None]
        return {
            "records": len(self.records),
            "brands": len(self.brands),
            "min_year": min(years) if years else None,
            "max_year": max(years) if years else None,
        }

    def filter(
        self,
        brand: str | None = None,
        keyword: str = "",
        min_year: int | None = None,
        max_year: int | None = None,
    ) -> list[dict]:
        normalized_brand = _compact(brand) if brand else ""
        normalized_keyword = _compact(keyword)
        selected = []
        for index, record in enumerate(self.records):
            if normalized_brand and self._brands[index] != normalized_brand:
                continue
            if normalized_keyword and normalized_keyword not in self._search_text[index]:
                continue
            year = self._years[index]
            if min_year is not None and (year is None or year < min_year):
                continue
            if max_year is not None and (year is None or year > max_year):
                continue
            selected.append(record)
        return selected

    def find_in_question(self, question: str, limit: int = 4) -> list[dict]:
        if limit <= 0 or not question.strip():
            return []
        text = _canonical(question).replace("(", "").replace(")", "")
        matches = [
            (match.start(), match.end(), indexes[0])
            for pattern, indexes in self._model_patterns
            for match in pattern.finditer(text)
        ]
        matches.sort(key=lambda match: (match[0], -(match[1] - match[0]), match[2]))
        selected = []
        spans = []
        seen_models = set()
        for start, end, index in matches:
            if any(start >= previous_start and end <= previous_end for previous_start, previous_end in spans):
                continue
            spans.append((start, end))
            record = self.records[index]
            model = _compact(str(record["产品型号"]))
            if model in seen_models:
                continue
            seen_models.add(model)
            selected.append(record)
            if len(selected) == limit:
                break
        return selected
